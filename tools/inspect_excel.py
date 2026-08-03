#!/usr/bin/env python3
"""Cartographie le classeur Excel du catalogue industriel.

Objectif : comprendre la structure réelle du classeur (onglets, listes déroulantes,
formules, colonnes masquées, grilles de croix) avant de migrer les données.

    python3 tools/inspect_excel.py "data/Catalogue industriel.xlsm" > docs/analyse-classeur.md

Le rapport produit est du Markdown, destiné à être relu et annoté avec l'IMI.
"""

import argparse
import re
import sys
from collections import Counter
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("openpyxl requis : pip install openpyxl")

# Marqueurs utilisés dans les grilles de croix du classeur.
CROSS_MARKS = {"x", "X", "✓", "✔", "1", "o", "O"}


def describe_sheet(ws, ws_values) -> None:
    print(f"\n## Onglet « {ws.title} »\n")
    print(f"- Dimensions : {ws.dimensions} ({ws.max_row} lignes × {ws.max_column} colonnes)")
    print(f"- État : {ws.sheet_state}")

    hidden_cols = [k for k, d in ws.column_dimensions.items() if d.hidden]
    hidden_rows = [k for k, d in ws.row_dimensions.items() if d.hidden]
    if hidden_cols:
        print(f"- **Colonnes masquées** : {', '.join(sorted(hidden_cols))}")
    if hidden_rows:
        print(f"- **Lignes masquées** : {len(hidden_rows)} ({hidden_rows[:20]}…)")
    if ws.merged_cells.ranges:
        print(f"- Cellules fusionnées : {len(ws.merged_cells.ranges)}")

    # Les listes déroulantes portent les options proposées au commercial.
    validations = list(getattr(ws, "data_validations", []).dataValidation or [])
    if validations:
        print(f"\n### Listes déroulantes ({len(validations)})\n")
        for dv in validations:
            print(f"- `{dv.sqref}` → type `{dv.type}`, source : `{dv.formula1}`")

    # Inventaire des formules : révèle les dépendances entre onglets.
    formulas: Counter = Counter()
    formula_examples: dict[str, str] = {}
    cross_cells: list[str] = []
    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            if isinstance(value, str) and value.startswith("="):
                head = re.match(r"=\s*([A-Z_.]+)\s*\(", value.upper())
                key = head.group(1) if head else "(référence directe)"
                formulas[key] += 1
                formula_examples.setdefault(key, f"{cell.coordinate} : {value[:160]}")
            elif isinstance(value, str) and value.strip() in CROSS_MARKS:
                cross_cells.append(cell.coordinate)

    if formulas:
        print(f"\n### Formules ({sum(formulas.values())} cellules)\n")
        for key, count in formulas.most_common():
            print(f"- `{key}` × {count} — ex. `{formula_examples[key]}`")

    if cross_cells:
        print(f"\n### Marqueurs de grille : {len(cross_cells)} cellules")
        print(f"  (ex. {', '.join(cross_cells[:15])}…)")

    # Aperçu des valeurs calculées, pour lire la grille telle que la voit l'utilisateur.
    print("\n### Aperçu (30 premières lignes, valeurs calculées)\n")
    print("```")
    for row in ws_values.iter_rows(min_row=1, max_row=30, max_col=min(ws.max_column, 20)):
        cells = ["" if c.value is None else str(c.value)[:22] for c in row]
        if any(cells):
            print(" | ".join(cells).rstrip(" |"))
    print("```")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()

    if not args.workbook.exists():
        sys.exit(f"Fichier introuvable : {args.workbook}")

    # Deux lectures : une avec les formules, une avec les dernières valeurs calculées.
    wb = load_workbook(args.workbook, data_only=False, keep_vba=True)
    wb_values = load_workbook(args.workbook, data_only=True, keep_vba=True)

    print(f"# Analyse du classeur — {args.workbook.name}\n")
    print(f"- Onglets : {len(wb.sheetnames)} — {', '.join(wb.sheetnames)}")
    hidden = [ws.title for ws in wb.worksheets if ws.sheet_state != "visible"]
    if hidden:
        print(f"- **Onglets masqués** : {', '.join(hidden)}")
    if wb.defined_names:
        print(f"\n## Noms définis ({len(wb.defined_names)})\n")
        for name, dn in wb.defined_names.items():
            print(f"- `{name}` → `{dn.value}`")

    for ws in wb.worksheets:
        describe_sheet(ws, wb_values[ws.title])


if __name__ == "__main__":
    main()
